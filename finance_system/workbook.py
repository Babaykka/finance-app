from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from finance_system.schema import LISTS, SHEETS, SheetSpec, sheet_by_name


DEFAULT_OUTPUT = Path("dist/finance-system-template.xlsx")
DATA_ROW_LIMIT = 1000

HEADER_FILL = PatternFill("solid", fgColor="EAF2F8")
HEADER_FONT = Font(bold=True)
TRUSTED_PERSON_FILL = PatternFill("solid", fgColor="FFF2CC")
CHECK_FILL = PatternFill("solid", fgColor="FCE4D6")
OVERDUE_FILL = PatternFill("solid", fgColor="F4CCCC")


def column_letter_for(sheet_name: str, column_key: str) -> str:
    sheet_spec = sheet_by_name(sheet_name)
    for column_index, column in enumerate(sheet_spec.columns, start=1):
        if column.key == column_key:
            return get_column_letter(column_index)
    raise KeyError(f"Unknown column: {sheet_name}.{column_key}")


def build_workbook(output_path: Path = DEFAULT_OUTPUT) -> Path:
    output_path = Path(output_path)
    workbook = Workbook()

    first_sheet = workbook.active
    first_sheet.title = SHEETS[0].name
    for sheet_spec in SHEETS[1:]:
        workbook.create_sheet(sheet_spec.name)

    for sheet_spec in SHEETS:
        worksheet = workbook[sheet_spec.name]
        if sheet_spec.name == "01 Главная":
            _write_dashboard(worksheet)
        elif sheet_spec.name == "06 Справочники":
            _write_reference_lists(worksheet)
        else:
            _write_data_sheet(worksheet, sheet_spec)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _write_dashboard(worksheet) -> None:
    storage_asset = column_letter_for("02 Места хранения", "asset")
    storage_amount = column_letter_for("02 Места хранения", "amount")
    storage_type = column_letter_for("02 Места хранения", "storage_type")
    storage_status = column_letter_for("02 Места хранения", "status")
    storage_priority = column_letter_for("02 Места хранения", "attention_priority")
    debt_asset = column_letter_for("03 Долги", "asset")
    debt_remaining_amount = column_letter_for("03 Долги", "remaining_amount")

    worksheet["A1"] = "Финансовая система"
    worksheet["A1"].font = Font(bold=True, size=16)
    worksheet["A4"] = "Актив"
    worksheet["B4"] = "Места хранения"
    worksheet["C4"] = "Долги"
    for row_index, asset in enumerate(LISTS["ВалютыАктивы"], start=5):
        worksheet.cell(row=row_index, column=1, value=asset)
        worksheet.cell(
            row=row_index,
            column=2,
            value=(
                f"=SUMIF('02 Места хранения'!${storage_asset}:${storage_asset},"
                f"A{row_index},'02 Места хранения'!${storage_amount}:${storage_amount})"
            ),
        )
        worksheet.cell(
            row=row_index,
            column=3,
            value=(
                f"=SUMIF('03 Долги'!${debt_asset}:${debt_asset},"
                f"A{row_index},'03 Долги'!${debt_remaining_amount}:${debt_remaining_amount})"
            ),
        )

    for column_index, header in enumerate(["Место", "Тип", "Валюта/актив", "Сумма", "Статус", "Приоритет"], start=8):
        worksheet.cell(row=4, column=column_index, value=header)
    worksheet["H5"] = (
        f'=FILTER({{\'02 Места хранения\'!B:B,\'02 Места хранения\'!{storage_type}:{storage_type},'
        f'\'02 Места хранения\'!{storage_asset}:{storage_asset},\'02 Места хранения\'!{storage_amount}:{storage_amount},'
        f'\'02 Места хранения\'!{storage_status}:{storage_status},\'02 Места хранения\'!{storage_priority}:{storage_priority}}},'
        f'(\'02 Места хранения\'!{storage_status}:{storage_status}="проверить")'
        f'+(\'02 Места хранения\'!{storage_type}:{storage_type}="у доверенного лица")'
        f'+(\'02 Места хранения\'!{storage_priority}:{storage_priority}="срочно"))'
    )
    _write_storage_distribution(worksheet)
    _write_payment_debt_and_decision_blocks(worksheet)

    for row in worksheet.iter_rows(min_row=4, max_row=4, min_col=1, max_col=20):
        for cell in row:
            if cell.value:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

    worksheet.freeze_panes = "A4"
    for column_index, width in {1: 16, 2: 20, 3: 18, 8: 24, 9: 28, 10: 18, 11: 16, 12: 18, 13: 18, 14: 30}.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def _write_storage_distribution(worksheet) -> None:
    worksheet["N3"] = "Распределение мест хранения"
    worksheet["N3"].font = HEADER_FONT
    worksheet["N4"] = "Тип места"
    for column_index, asset in enumerate(LISTS["ВалютыАктивы"], start=15):
        worksheet.cell(row=4, column=column_index, value=asset)

    for row_index, storage_type in enumerate(LISTS["ТипыМест"], start=5):
        worksheet.cell(row=row_index, column=14, value=storage_type)
        for column_index in range(15, 15 + len(LISTS["ВалютыАктивы"])):
            asset_cell = f"{get_column_letter(column_index)}$4"
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=(
                    "=SUMIFS('02 Места хранения'!$G:$G,"
                    f"'02 Места хранения'!$C:$C,$N{row_index},"
                    f"'02 Места хранения'!$F:$F,{asset_cell})"
                ),
            )


def _write_payment_debt_and_decision_blocks(worksheet) -> None:
    worksheet["A14"] = "Ближайшие ожидаемые платежи"
    worksheet["A15"] = "Дата"
    worksheet["B15"] = "Должник/клиент"
    worksheet["C15"] = "Сумма"
    worksheet["D15"] = "Статус"
    worksheet["A16"] = (
        '=SORT(FILTER({\'04 График и оплаты\'!E:E,\'04 График и оплаты\'!C:C,'
        '\'04 График и оплаты\'!F:F,\'04 График и оплаты\'!H:H},'
        '\'04 График и оплаты\'!D:D="план",\'04 График и оплаты\'!H:H<>"отменено"),1,TRUE)'
    )

    worksheet["E14"] = "Статусы долгов"
    worksheet["E15"] = "Статус"
    worksheet["F15"] = "Количество"
    for row_index, status in enumerate(LISTS["СтатусыДолгов"], start=16):
        worksheet.cell(row=row_index, column=5, value=status)
        worksheet.cell(row=row_index, column=6, value=f'=COUNTIF(\'03 Долги\'!$H:$H,E{row_index})')

    worksheet["H14"] = "Решения"
    worksheet["H15"] = "Дата"
    worksheet["I15"] = "Область"
    worksheet["J15"] = "Действие"
    worksheet["K15"] = "Приоритет"
    worksheet["L15"] = "Статус"
    worksheet["H16"] = (
        '=FILTER({\'08 Решения\'!B:B,\'08 Решения\'!C:C,\'08 Решения\'!D:D,'
        '\'08 Решения\'!G:G,\'08 Решения\'!H:H},'
        '\'08 Решения\'!H:H<>"сделано",\'08 Решения\'!H:H<>"")'
    )


def _write_reference_lists(worksheet) -> None:
    for column_index, (list_name, values) in enumerate(LISTS.items(), start=1):
        column_letter = get_column_letter(column_index)
        worksheet.cell(row=1, column=column_index, value=list_name)
        worksheet.cell(row=1, column=column_index).fill = HEADER_FILL
        worksheet.cell(row=1, column=column_index).font = HEADER_FONT
        worksheet.column_dimensions[column_letter].width = max(18, len(list_name) + 2)
        for row_index, value in enumerate(values, start=2):
            worksheet.cell(row=row_index, column=column_index, value=value)

    worksheet.freeze_panes = "A2"


def _write_data_sheet(worksheet, sheet_spec: SheetSpec) -> None:
    for column_index, column in enumerate(sheet_spec.columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column.title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        worksheet.column_dimensions[get_column_letter(column_index)].width = column.width

        if column.kind == "number":
            for row_index in range(2, DATA_ROW_LIMIT + 1):
                worksheet.cell(row=row_index, column=column_index).number_format = "#,##0.00"
        elif column.kind == "date":
            for row_index in range(2, DATA_ROW_LIMIT + 1):
                worksheet.cell(row=row_index, column=column_index).number_format = "yyyy-mm-dd"

        if column.dropdown:
            _add_dropdown(worksheet, column_index, column.dropdown)

    if sheet_spec.name == "03 Долги":
        _write_debt_remaining_formulas(worksheet)

    worksheet.freeze_panes = f"A{sheet_spec.frozen_rows + 1}"
    worksheet.auto_filter.ref = worksheet.dimensions
    _add_conditional_formatting(worksheet, sheet_spec)


def _write_debt_remaining_formulas(worksheet) -> None:
    debt_id = column_letter_for("03 Долги", "id_debt")
    original_amount = column_letter_for("03 Долги", "original_amount")
    remaining_amount = column_letter_for("03 Долги", "remaining_amount")
    payment_debt_id = column_letter_for("04 График и оплаты", "id_debt")
    payment_record_type = column_letter_for("04 График и оплаты", "record_type")
    payment_amount = column_letter_for("04 График и оплаты", "amount")
    payment_status = column_letter_for("04 График и оплаты", "status")

    for row_index in range(2, DATA_ROW_LIMIT + 1):
        worksheet[f"{remaining_amount}{row_index}"] = (
            f'=IF({debt_id}{row_index}="","",{original_amount}{row_index}-'
            f"SUMIFS('04 График и оплаты'!${payment_amount}:${payment_amount},"
            f"'04 График и оплаты'!${payment_debt_id}:${payment_debt_id},{debt_id}{row_index},"
            f'\'04 График и оплаты\'!${payment_record_type}:${payment_record_type},"факт",'
            f'\'04 График и оплаты\'!${payment_status}:${payment_status},"<>отменено"))'
        )


def _add_dropdown(worksheet, column_index: int, list_name: str) -> None:
    reference_column = list(LISTS).index(list_name) + 1
    reference_letter = get_column_letter(reference_column)
    last_row = len(LISTS[list_name]) + 1
    formula = f"'06 Справочники'!${reference_letter}$2:${reference_letter}${last_row}"
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Выберите значение из справочника."
    validation.errorTitle = "Недопустимое значение"
    worksheet.add_data_validation(validation)
    column_letter = get_column_letter(column_index)
    validation.add(f"{column_letter}2:{column_letter}{DATA_ROW_LIMIT}")


def _add_conditional_formatting(worksheet, sheet_spec: SheetSpec) -> None:
    last_column = get_column_letter(len(sheet_spec.columns))
    full_range = f"A2:{last_column}{DATA_ROW_LIMIT}"

    if sheet_spec.name == "02 Места хранения":
        storage_type = column_letter_for(sheet_spec.name, "storage_type")
        checked_at = column_letter_for(sheet_spec.name, "checked_at")
        review_or_end_date = column_letter_for(sheet_spec.name, "review_or_end_date")
        storage_status = column_letter_for(sheet_spec.name, "status")
        worksheet.conditional_formatting.add(
            full_range,
            FormulaRule(formula=[f'${storage_type}2="у доверенного лица"'], fill=TRUSTED_PERSON_FILL),
        )
        worksheet.conditional_formatting.add(
            full_range,
            FormulaRule(formula=[f'${storage_status}2="проверить"'], fill=CHECK_FILL),
        )
        worksheet.conditional_formatting.add(
            full_range,
            FormulaRule(
                formula=[f'AND($A2<>"",${checked_at}2<TODAY()-30,${storage_status}2<>"закрыто")'],
                fill=CHECK_FILL,
            ),
        )
        worksheet.conditional_formatting.add(
            full_range,
            FormulaRule(
                formula=[
                    f'AND($A2<>"",${storage_type}2="доходное банковское хранение",${review_or_end_date}2<>"",'
                    f'${review_or_end_date}2<=TODAY()+14,${storage_status}2<>"закрыто")'
                ],
                fill=TRUSTED_PERSON_FILL,
            ),
        )
    elif sheet_spec.name == "03 Долги":
        debt_status = column_letter_for(sheet_spec.name, "status")
        worksheet.conditional_formatting.add(
            full_range,
            FormulaRule(formula=[f'${debt_status}2="просрочено"'], fill=OVERDUE_FILL),
        )


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the finance system workbook template.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    output_path = build_workbook(args.output)
    print(f"Workbook written to {output_path}")


if __name__ == "__main__":
    main()
