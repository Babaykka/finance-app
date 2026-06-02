import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from finance_system.schema import LISTS, SHEETS
from finance_system.workbook import DATA_ROW_LIMIT, build_workbook, column_letter_for


BUNDLED_PYTHON = Path("/Users/hrach/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3")


class WorkbookTest(unittest.TestCase):
    def test_builds_expected_sheets_and_headers(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        self.assertEqual(
            workbook.sheetnames,
            [
                "01 Главная",
                "02 Места хранения",
                "03 Долги",
                "04 График и оплаты",
                "05 Крупные движения",
                "06 Справочники",
                "07 Расходы-план",
                "08 Решения",
            ],
        )
        storage = workbook["02 Места хранения"]
        self.assertEqual(storage["A1"].value, "ID места хранения")
        self.assertEqual(storage["C1"].value, "Тип места")
        self.assertEqual(storage["J1"].value, "Правило расчета процентов")
        self.assertEqual(storage["L1"].value, "День выплаты процентов")

    def test_reference_lists_are_written(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        refs = workbook["06 Справочники"]
        headers = [cell.value for cell in refs[1] if cell.value]
        self.assertIn("ТипыМест", headers)
        self.assertIn("ПодтипыМест", headers)
        self.assertIn("ПравилаПроцентов", headers)
        self.assertIn("ВлияниеСнятия", headers)
        self.assertIn("доходное банковское хранение", [refs.cell(row=i, column=2).value for i in range(2, 10)])
        self.assertEqual(len(headers), len(LISTS))

    def test_dashboard_contains_google_sheets_formulas(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        dashboard = workbook["01 Главная"]
        self.assertEqual(dashboard["A1"].value, "Финансовая система")
        self.assertEqual(
            [dashboard.cell(row=row, column=1).value for row in range(5, 5 + len(LISTS["ВалютыАктивы"]))],
            LISTS["ВалютыАктивы"],
        )
        self.assertEqual(dashboard["B5"].value, "=SUMIF('02 Места хранения'!$F:$F,A5,'02 Места хранения'!$G:$G)")
        self.assertEqual(dashboard["C5"].value, "=SUMIF('03 Долги'!$E:$E,A5,'03 Долги'!$G:$G)")
        self.assertEqual(
            [dashboard.cell(row=4, column=column).value for column in range(8, 14)],
            ["Место", "Тип", "Валюта/актив", "Сумма", "Статус", "Приоритет"],
        )
        self.assertIn("FILTER", dashboard["H5"].value)
        self.assertIn("CHOOSECOLS", dashboard["H5"].value)
        self.assertIn("проверить", dashboard["H5"].value)
        self.assertIn("у доверенного лица", dashboard["H5"].value)
        self.assertIn("срочно", dashboard["H5"].value)

    def test_dashboard_contains_operational_blocks(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        dashboard = workbook["01 Главная"]
        self.assertEqual(dashboard["N3"].value, "Распределение мест хранения")
        self.assertEqual([dashboard.cell(row=4, column=column).value for column in range(15, 15 + len(LISTS["ВалютыАктивы"]))], LISTS["ВалютыАктивы"])
        self.assertIn("SUMIFS", dashboard["O5"].value)
        self.assertIn("'02 Места хранения'", dashboard["O5"].value)
        self.assertIn("$C:$C", dashboard["O5"].value)
        self.assertIn("$F:$F", dashboard["O5"].value)

        self.assertEqual(dashboard["A14"].value, "Ближайшие ожидаемые платежи")
        self.assertIn("FILTER", dashboard["A16"].value)
        self.assertIn("CHOOSECOLS", dashboard["A16"].value)
        self.assertIn("'04 График и оплаты'", dashboard["A16"].value)
        self.assertEqual(dashboard["E14"].value, "Статусы долгов")
        self.assertIn("COUNTIF", dashboard["F16"].value)
        self.assertEqual(dashboard["H14"].value, "Решения")
        self.assertIn("FILTER", dashboard["H16"].value)
        self.assertIn("CHOOSECOLS", dashboard["H16"].value)
        self.assertIn("'08 Решения'", dashboard["H16"].value)

    def test_formulas_do_not_use_google_sheets_range_array_literals(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        formulas = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas.append((worksheet.title, cell.coordinate, cell.value))

        offenders = [
            (sheet, coordinate, formula)
            for sheet, coordinate, formula in formulas
            if "{" in formula and "!" in formula
        ]
        self.assertEqual(offenders, [])

    def test_debt_remaining_amount_is_formula_derived_from_actual_payments(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        formula = workbook["03 Долги"]["G2"].value
        self.assertEqual(
            formula,
            '=IF(A2="","",F2-SUMIFS(\'04 График и оплаты\'!$F:$F,\'04 График и оплаты\'!$B:$B,A2,\'04 График и оплаты\'!$D:$D,"факт",\'04 График и оплаты\'!$H:$H,"<>отменено"))',
        )

    def test_data_validations_are_present(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        storage = workbook["02 Места хранения"]
        self.assertGreaterEqual(storage.data_validations.count, 8)

    def test_column_letter_for_uses_schema_columns(self):
        self.assertEqual(column_letter_for("02 Места хранения", "asset"), "F")
        self.assertEqual(column_letter_for("02 Места хранения", "amount"), "G")
        self.assertEqual(column_letter_for("03 Долги", "status"), "H")

    def test_dropdown_validations_cover_all_schema_dropdowns(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        for sheet_spec in SHEETS:
            dropdown_columns = [column for column in sheet_spec.columns if column.dropdown]
            if not dropdown_columns:
                continue

            worksheet = workbook[sheet_spec.name]
            validations = list(worksheet.data_validations.dataValidation)
            self.assertGreaterEqual(len(validations), len(dropdown_columns), sheet_spec.name)
            validations_by_range = {str(validation.sqref): validation for validation in validations}
            for column in dropdown_columns:
                column_letter = column_letter_for(sheet_spec.name, column.key)
                range_ref = f"{column_letter}2:{column_letter}{DATA_ROW_LIMIT}"
                list_column = get_column_letter(list(LISTS).index(column.dropdown) + 1)
                expected_formula = f"'06 Справочники'!${list_column}$2:${list_column}${len(LISTS[column.dropdown]) + 1}"

                self.assertIn(range_ref, validations_by_range, f"{sheet_spec.name}.{column.key}")
                self.assertEqual(validations_by_range[range_ref].formula1, expected_formula)

    def test_cli_writes_workbook(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "cli-template.xlsx"
            result = subprocess.run(
                [str(BUNDLED_PYTHON), "-m", "finance_system.workbook", "--output", str(output)],
                check=False,
                capture_output=True,
                cwd=Path(__file__).resolve().parents[1],
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(result.stdout.strip(), f"Workbook written to {output}")
            self.assertTrue(output.exists())

    def test_conditional_formatting_rules_are_present(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "finance-system-template.xlsx"
            build_workbook(output)
            workbook = load_workbook(output, data_only=False)

        storage_formulas = _conditional_formatting_formulas(workbook["02 Места хранения"])
        debt_formulas = _conditional_formatting_formulas(workbook["03 Долги"])

        self.assertIn('$C2="у доверенного лица"', storage_formulas)
        self.assertIn('$Q2="проверить"', storage_formulas)
        stale_checked_formula = _formula_containing(storage_formulas, "TODAY()-30")
        self.assertIn('$A2<>""', stale_checked_formula)
        self.assertIn("TODAY()-30", stale_checked_formula)
        income_review_formula = _formula_containing(storage_formulas, "TODAY()+14")
        self.assertIn('$A2<>""', income_review_formula)
        self.assertIn('$O2<>""', income_review_formula)
        self.assertIn("TODAY()+14", income_review_formula)
        self.assertIn("доходное банковское хранение", income_review_formula)
        self.assertIn('$H2="просрочено"', debt_formulas)


def _conditional_formatting_formulas(worksheet):
    formulas = []
    for conditional_formatting in worksheet.conditional_formatting:
        for rule in conditional_formatting.rules:
            formulas.extend(rule.formula)
    return formulas


def _formula_containing(formulas, expected_text):
    for formula in formulas:
        if expected_text in formula:
            return formula
    raise AssertionError(f"No conditional formatting formula contains {expected_text!r}: {formulas}")


if __name__ == "__main__":
    unittest.main()
